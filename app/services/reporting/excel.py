"""Professional Excel workbook generation."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.config import get_settings
from app.core.exceptions import ExportError
from app.db.models.opportunity import Opportunity
from app.services.normalisation.dates import utcnow

HEADER_FILL = PatternFill("solid", fgColor="1B3A4B")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri")
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1B3A4B")
THIN = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)
ZEBRA = PatternFill("solid", fgColor="F6F8FA")


def _header(ws: Worksheet, headers: list[str], row: int = 1) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"
    ws.freeze_panes = f"A{row + 1}"
    ws.row_dimensions[row].height = 22


def _autosize(ws: Worksheet, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_rows(ws: Worksheet, rows: list[list[Any]], start: int = 2) -> None:
    for offset, row in enumerate(rows):
        excel_row = start + offset
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=excel_row, column=col, value=_excel_value(value))
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=col in {1, 5})
            if offset % 2:
                cell.fill = ZEBRA
            if isinstance(value, datetime):
                cell.number_format = "YYYY-MM-DD HH:MM"
            if isinstance(value, Decimal):
                cell.number_format = "£#,##0.00"


def _stringify_dt(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is not None:
        return value.astimezone().replace(tzinfo=None).isoformat()
    return value.isoformat()


def _excel_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone().replace(tzinfo=None)
    return value


def _score_scale(ws: Worksheet, score_col: str, start: int, end: int) -> None:
    if end < start:
        return
    ws.conditional_formatting.add(
        f"{score_col}{start}:{score_col}{end}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="F4C7C3",
            mid_type="num",
            mid_value=50,
            mid_color="FFE599",
            end_type="num",
            end_value=100,
            end_color="B7E1CD",
        ),
    )


class ExcelReportBuilder:
    def build(
        self,
        opportunities: Sequence[Opportunity],
        metrics: dict[str, Any],
        *,
        run_id: UUID | None = None,
        path: Path | None = None,
    ) -> Path:
        settings = get_settings()
        settings.ensure_data_dirs()
        stamp = utcnow().strftime("%Y%m%d_%H%M%S")
        target = path or (settings.exports_dir / f"opportunities_{stamp}.xlsx")
        try:
            wb = Workbook()
            self._summary(wb.active, opportunities, metrics, run_id)
            self._opportunity_sheet(
                wb.create_sheet("Top Opportunities"),
                [item for item in opportunities if item.score >= 65][:50],
            )
            self._opportunity_sheet(wb.create_sheet("All Opportunities"), list(opportunities))
            self._opportunity_sheet(
                wb.create_sheet("Closing Soon"),
                [
                    item
                    for item in opportunities
                    if item.deadline_status in {"CRITICAL", "URGENT", "UPCOMING"}
                ],
            )
            self._opportunity_sheet(
                wb.create_sheet("New Since Last Run"),
                [
                    item
                    for item in opportunities
                    if metrics.get("new_ids") and str(item.id) in set(metrics.get("new_ids") or [])
                ]
                or sorted(
                    opportunities,
                    key=lambda item: _stringify_dt(item.first_seen_at) or "",
                    reverse=True,
                )[:40],
            )
            self._opportunity_sheet(
                wb.create_sheet("Changed Opportunities"),
                [
                    item
                    for item in opportunities
                    if _stringify_dt(item.created_at) != _stringify_dt(item.updated_at)
                ][:80],
            )
            self._opportunity_sheet(
                wb.create_sheet("Rejected Low Score"),
                [item for item in opportunities if item.score < 25],
            )
            self._metrics_sheet(wb.create_sheet("Pipeline Metrics"), metrics)
            self._source_sheet(wb.create_sheet("Source Performance"), opportunities)
            wb.save(target)
            return target
        except Exception as exc:
            raise ExportError(f"Failed to write Excel report: {exc}") from exc

    def _summary(
        self,
        ws: Worksheet,
        opportunities: Sequence[Opportunity],
        metrics: dict[str, Any],
        run_id: UUID | None,
    ) -> None:
        ws.title = "Executive Summary"
        ws["A1"] = "Opportunity Intelligence Automation Engine"
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:D1")
        ws["A2"] = "Run summary"
        ws["A3"] = "Run ID"
        ws["B3"] = str(run_id or "")
        rows = [
            ("Generated", utcnow().isoformat()),
            ("Total opportunities", len(opportunities)),
            ("Discovered this run", metrics.get("records_discovered", "")),
            ("New", metrics.get("new_records", "")),
            ("Updated", metrics.get("updated_records", "")),
            ("Duplicates", metrics.get("duplicates", "")),
            ("Invalid", metrics.get("invalid_records", "")),
            ("High priority", metrics.get("high_priority_opportunities", "")),
            ("Duration (seconds)", metrics.get("duration_seconds", "")),
            ("Sources attempted", metrics.get("sources_attempted", "")),
            ("Sources successful", metrics.get("sources_successful", "")),
            ("Sources failed", metrics.get("sources_failed", "")),
        ]
        for index, (label, value) in enumerate(rows, start=4):
            ws.cell(row=index, column=1, value=label).font = Font(bold=True)
            ws.cell(row=index, column=2, value=value)
        _autosize(ws, {1: 32, 2: 48, 3: 24, 4: 24})

    def _opportunity_sheet(self, ws: Worksheet, items: Sequence[Opportunity]) -> None:
        headers = [
            "Title",
            "Organisation",
            "Score",
            "Rating",
            "Deadline",
            "Deadline status",
            "Value",
            "Currency",
            "Location",
            "Category",
            "Status",
            "Source URL",
            "First seen",
            "Reasons",
        ]
        _header(ws, headers)
        rows = [
            [
                item.title,
                item.organisation,
                item.score,
                item.score_rating,
                item.deadline_at,
                item.deadline_status,
                item.estimated_value,
                item.currency,
                item.location,
                item.category,
                item.status,
                item.source_url,
                item.first_seen_at,
                "; ".join(item.score_reasons or []),
            ]
            for item in items
        ]
        _write_rows(ws, rows)
        for index, item in enumerate(items, start=2):
            if item.source_url:
                cell = ws.cell(row=index, column=12)
                cell.hyperlink = item.source_url
                cell.font = Font(color="1155CC", underline="single")
        last = len(items) + 1
        _score_scale(ws, "C", 2, last)
        if last >= 2:
            ws.conditional_formatting.add(
                f"F2:F{last}",
                FormulaRule(formula=['$F2="EXPIRED"'], fill=PatternFill("solid", fgColor="F4C7C3")),
            )
            ws.conditional_formatting.add(
                f"F2:F{last}",
                FormulaRule(
                    formula=['$F2="CRITICAL"'], fill=PatternFill("solid", fgColor="FCE8B2")
                ),
            )
        _autosize(
            ws,
            {
                1: 42,
                2: 32,
                3: 10,
                4: 12,
                5: 20,
                6: 16,
                7: 14,
                8: 10,
                9: 22,
                10: 18,
                11: 12,
                12: 36,
                13: 20,
                14: 50,
            },
        )

    def _metrics_sheet(self, ws: Worksheet, metrics: dict[str, Any]) -> None:
        _header(ws, ["Metric", "Value"])
        rows = [[key, value] for key, value in metrics.items()]
        _write_rows(ws, rows)
        _autosize(ws, {1: 36, 2: 40})

    def _source_sheet(self, ws: Worksheet, opportunities: Sequence[Opportunity]) -> None:
        _header(ws, ["Source ID", "Opportunities", "Average score", "High priority"])
        grouped: dict[str, list[Opportunity]] = {}
        for item in opportunities:
            grouped.setdefault(str(item.source_id), []).append(item)
        rows = []
        for source_id, items in grouped.items():
            high = len([item for item in items if item.score >= 65])
            avg = round(sum(item.score for item in items) / len(items), 2)
            rows.append([source_id, len(items), avg, high])
        _write_rows(ws, rows)
        _autosize(ws, {1: 40, 2: 18, 3: 18, 4: 16})
